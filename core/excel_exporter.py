"""
Multi-Tab Professional Excel BOQ and Audit Exporter using OpenPyXL.
Generates CPWD / L&T compliant Bill of Quantities workbooks with automated formatting.
"""

import io
from typing import Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class BOQExcelExporter:
    """
    Exports parsed BIM takeoff, aggregated BOQ, and anomaly audit records
    into a polished, multi-tab Microsoft Excel spreadsheet (.xlsx).
    """

    def __init__(self):
        # Color Palette: Corporate Civil Engineering
        self.navy_header_fill = PatternFill(start_color="1E3E62", end_color="1E3E62", fill_type="solid")
        self.sub_header_fill = PatternFill(start_color="2B4865", end_color="2B4865", fill_type="solid")
        self.accent_gold_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        self.high_risk_fill = PatternFill(start_color="FFD6D6", end_color="FFD6D6", fill_type="solid")
        self.med_risk_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        self.total_row_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")

        self.font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.font_title = Font(name="Calibri", size=15, bold=True, color="1E3E62")
        self.font_subtitle = Font(name="Calibri", size=10, italic=True, color="555555")
        self.font_bold = Font(name="Calibri", size=11, bold=True)
        self.font_regular = Font(name="Calibri", size=10)

        thin_side = Side(border_style="thin", color="CCCCCC")
        self.cell_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        self.total_border = Border(
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="double", color="000000")
        )

    def export_to_bytes(
        self,
        metadata: Dict[str, Any],
        kpis: Dict[str, Any],
        df_boq: pd.DataFrame,
        df_audited: pd.DataFrame,
        audit_summary: Dict[str, Any],
    ) -> bytes:
        """
        Builds the entire workbook in memory and returns bytes for Streamlit download.
        """
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Tab 1: Executive Summary
        self._build_summary_sheet(wb, metadata, kpis, audit_summary)

        # Tab 2: Abstract of Cost (BOQ)
        self._build_boq_sheet(wb, df_boq)

        # Tab 3: Detailed Itemized Takeoff
        self._build_takeoff_sheet(wb, df_audited)

        # Tab 4: AI Anomaly & Variance Audit
        self._build_anomaly_sheet(wb, df_audited)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _auto_adjust_columns(self, ws):
        """Auto fits column widths with comfortable padding."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val is not None:
                    # Ignore very long cells in title rows to avoid blowout
                    if cell.row in [1, 2, 3] and "OpenBIM" in str(val):
                        continue
                    length = len(str(val))
                    if length > max_len:
                        max_len = length
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

    def _build_summary_sheet(self, wb, metadata, kpis, audit_summary):
        ws = wb.create_sheet(title="Executive Summary")
        ws.views.sheetView[0].showGridLines = True

        ws["B2"] = "OpenBIM Automated BOQ & Cost Variance Intelligence Engine"
        ws["B2"].font = self.font_title
        ws["B3"] = "Executive Project Takeoff & Risk Verification Summary"
        ws["B3"].font = self.font_subtitle

        # Metadata Section
        ws["B5"] = "Project Metadata"
        ws["B5"].font = self.font_bold
        meta_items = [
            ("Project Name", metadata.get("project_name", "N/A")),
            ("Building / Block", metadata.get("building_name", "N/A")),
            ("IFC Schema", metadata.get("schema", "IFC4")),
            ("Total Model Entities", metadata.get("total_entities", 0)),
            ("Exporting BIM Application", metadata.get("originating_system", "N/A")),
        ]
        for idx, (label, val) in enumerate(meta_items, start=6):
            ws[f"B{idx}"] = label
            ws[f"B{idx}"].font = self.font_bold
            ws[f"C{idx}"] = str(val)
            ws[f"C{idx}"].font = self.font_regular

        # KPI Metrics Table
        ws["E5"] = "Key Takeoff Indicators"
        ws["E5"].font = self.font_bold
        kpi_items = [
            ("Total Structural Members", f"{kpis.get('total_elements', 0):,d}"),
            ("Total Concrete Volume", f"{kpis.get('total_concrete_cum', 0.0):,.2f} m³"),
            ("Total Reinforcement Steel", f"{kpis.get('total_steel_tonne', 0.0):,.2f} MT"),
            ("Overall Rebar Density", f"{kpis.get('rebar_density_overall', 0.0):,.1f} kg/m³"),
            ("Total Formwork / Shuttering", f"{kpis.get('total_formwork_sqm', 0.0):,.2f} m²"),
            ("Total Project Cost (INR)", f"₹{kpis.get('total_cost_inr', 0.0):,.2f}"),
            ("Total Cost in Lakhs", f"₹{kpis.get('total_cost_lakhs', 0.0):,.2f} Lakhs"),
        ]
        for idx, (label, val) in enumerate(kpi_items, start=6):
            ws[f"E{idx}"] = label
            ws[f"E{idx}"].font = self.font_bold
            ws[f"F{idx}"] = val
            ws[f"F{idx}"].font = self.font_regular

        # Anomaly Audit Card
        ws["B14"] = "AI Anomaly & Variance Audit Status"
        ws["B14"].font = self.font_bold
        audit_items = [
            ("High Risk Anomaly Elements", f"{audit_summary.get('high_risk_count', 0)} members"),
            ("Medium Risk Warning Elements", f"{audit_summary.get('medium_risk_count', 0)} members"),
            ("Verified Compliant Elements", f"{audit_summary.get('normal_count', 0)} members"),
            ("Total Cost at Risk (INR)", f"₹{audit_summary.get('at_risk_cost_inr', 0.0):,.2f}"),
            ("At Risk Cost in Lakhs", f"₹{audit_summary.get('at_risk_cost_lakhs', 0.0):,.2f} Lakhs"),
        ]
        for idx, (label, val) in enumerate(audit_items, start=15):
            ws[f"B{idx}"] = label
            ws[f"B{idx}"].font = self.font_bold
            ws[f"C{idx}"] = val
            ws[f"C{idx}"].font = self.font_regular
            if "High Risk" in label and audit_summary.get('high_risk_count', 0) > 0:
                ws[f"C{idx}"].fill = self.high_risk_fill

        self._auto_adjust_columns(ws)

    def _build_boq_sheet(self, wb, df_boq):
        ws = wb.create_sheet(title="Abstract of Cost (BOQ)")
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = "SCHEDULE OF QUANTITIES & ABSTRACT OF COST"
        ws["A1"].font = self.font_title
        ws["A2"] = "Aligned with CPWD Delhi Schedule of Rates (DSR 2023)"
        ws["A2"].font = self.font_subtitle

        headers = ["Item No.", "Work Category", "DSR Item Code", "Item Description", "Unit", "Quantity", "Rate (INR)", "Amount (INR)", "% Share"]
        start_row = 4

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = self.font_header
            cell.fill = self.navy_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        current_row = start_row + 1
        for i, (_, row) in enumerate(df_boq.iterrows(), start=1):
            ws.cell(row=current_row, column=1, value=i).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=2, value=row.get("Category", ""))
            ws.cell(row=current_row, column=3, value=row.get("Item Code", "")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=4, value=row.get("Description", ""))
            ws.cell(row=current_row, column=5, value=row.get("Unit", "")).alignment = Alignment(horizontal="center")
            
            c_qty = ws.cell(row=current_row, column=6, value=float(row.get("Quantity", 0)))
            c_qty.number_format = "#,##0.00"
            
            c_rate = ws.cell(row=current_row, column=7, value=float(row.get("Rate (INR)", 0)))
            c_rate.number_format = "#,##0.00"
            
            c_amt = ws.cell(row=current_row, column=8, value=float(row.get("Amount (INR)", 0)))
            c_amt.number_format = "#,##0.00"
            
            c_pct = ws.cell(row=current_row, column=9, value=float(row.get("% Share", 0)) / 100.0)
            c_pct.number_format = "0.00%"

            for col_idx in range(1, 10):
                ws.cell(row=current_row, column=col_idx).border = self.cell_border
                ws.cell(row=current_row, column=col_idx).font = self.font_regular
            current_row += 1

        # Total Row
        ws.cell(row=current_row, column=4, value="TOTAL ABSTRACT OF COST (INR)").font = self.font_bold
        ws.cell(row=current_row, column=4).alignment = Alignment(horizontal="right")
        total_cell = ws.cell(row=current_row, column=8, value=f"=SUM(H{start_row+1}:H{current_row-1})")
        total_cell.font = self.font_bold
        total_cell.number_format = "#,##0.00"

        for col_idx in range(1, 10):
            c = ws.cell(row=current_row, column=col_idx)
            c.border = self.total_border
            c.fill = self.total_row_fill

        self._auto_adjust_columns(ws)

    def _build_takeoff_sheet(self, wb, df_audited):
        ws = wb.create_sheet(title="Detailed Takeoff (M-Book)")
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = "ITEMIZED BIM MEASUREMENT SHEET (M-BOOK)"
        ws["A1"].font = self.font_title

        headers = [
            "Tag / ID", "Element Name", "IFC Type", "Storey / Level", "Concrete Grade",
            "Net Volume (m³)", "Rebar (kg)", "Rebar (MT)", "Formwork (m²)",
            "Concrete Cost (₹)", "Steel Cost (₹)", "Formwork Cost (₹)", "Total Cost (₹)", "Risk Level"
        ]
        start_row = 3

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = self.font_header
            cell.fill = self.sub_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        current_row = start_row + 1
        for _, row in df_audited.iterrows():
            ws.cell(row=current_row, column=1, value=str(row.get("tag", "")))
            ws.cell(row=current_row, column=2, value=str(row.get("name", "")))
            ws.cell(row=current_row, column=3, value=str(row.get("ifc_type", "")).replace("Ifc", ""))
            ws.cell(row=current_row, column=4, value=str(row.get("storey", "")))
            ws.cell(row=current_row, column=5, value=str(row.get("concrete_grade", "")))
            
            c_vol = ws.cell(row=current_row, column=6, value=float(row.get("concrete_vol_cum", 0)))
            c_vol.number_format = "#,##0.000"
            
            c_rkg = ws.cell(row=current_row, column=7, value=float(row.get("rebar_kg", 0)))
            c_rkg.number_format = "#,##0.0"
            
            c_rmt = ws.cell(row=current_row, column=8, value=float(row.get("rebar_tonne", 0)))
            c_rmt.number_format = "#,##0.000"
            
            c_fsq = ws.cell(row=current_row, column=9, value=float(row.get("formwork_sqm", 0)))
            c_fsq.number_format = "#,##0.00"

            c_cc = ws.cell(row=current_row, column=10, value=float(row.get("concrete_cost_inr", 0)))
            c_cc.number_format = "#,##0.00"
            
            c_sc = ws.cell(row=current_row, column=11, value=float(row.get("steel_cost_inr", 0)))
            c_sc.number_format = "#,##0.00"
            
            c_fc = ws.cell(row=current_row, column=12, value=float(row.get("formwork_cost_inr", 0)))
            c_fc.number_format = "#,##0.00"
            
            c_tc = ws.cell(row=current_row, column=13, value=float(row.get("total_cost_inr", 0)))
            c_tc.number_format = "#,##0.00"

            c_risk = ws.cell(row=current_row, column=14, value=str(row.get("risk_level", "NORMAL")))
            c_risk.alignment = Alignment(horizontal="center")
            if row.get("risk_level") == "HIGH":
                c_risk.fill = self.high_risk_fill
            elif row.get("risk_level") == "MEDIUM":
                c_risk.fill = self.med_risk_fill

            for col_idx in range(1, 15):
                ws.cell(row=current_row, column=col_idx).border = self.cell_border
                ws.cell(row=current_row, column=col_idx).font = self.font_regular
            current_row += 1

        self._auto_adjust_columns(ws)

    def _build_anomaly_sheet(self, wb, df_audited):
        ws = wb.create_sheet(title="Cost Anomaly Audit")
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = "AI COST & QUANTITY VARIANCE AUDIT LOG"
        ws["A1"].font = self.font_title
        ws["A2"] = "Isolation Forest & Cohort Z-Score Outlier Verification"
        ws["A2"].font = self.font_subtitle

        headers = [
            "Element Name", "IFC Type", "Storey", "Net Volume (m³)", "Peer Mean (m³)",
            "Z-Score", "Cost (INR)", "Risk Level", "Audit Diagnosis & Reasons", "Recommended QS Action"
        ]
        start_row = 4

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = self.font_header
            cell.fill = self.navy_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        anomalies_df = df_audited[df_audited["risk_level"].isin(["HIGH", "MEDIUM"])].copy()
        current_row = start_row + 1

        for _, row in anomalies_df.iterrows():
            ws.cell(row=current_row, column=1, value=str(row.get("name", "")))
            ws.cell(row=current_row, column=2, value=str(row.get("ifc_type", "")).replace("Ifc", ""))
            ws.cell(row=current_row, column=3, value=str(row.get("storey", "")))
            
            c_v = ws.cell(row=current_row, column=4, value=float(row.get("concrete_vol_cum", 0)))
            c_v.number_format = "#,##0.000"
            
            c_m = ws.cell(row=current_row, column=5, value=float(row.get("cohort_mean_vol", 0)))
            c_m.number_format = "#,##0.000"
            
            c_z = ws.cell(row=current_row, column=6, value=float(row.get("z_score_vol", 0)))
            c_z.number_format = "0.00"
            
            c_cost = ws.cell(row=current_row, column=7, value=float(row.get("total_cost_inr", 0)))
            c_cost.number_format = "#,##0.00"
            
            c_risk = ws.cell(row=current_row, column=8, value=str(row.get("risk_level", "")))
            c_risk.alignment = Alignment(horizontal="center")
            c_risk.fill = self.high_risk_fill if row.get("risk_level") == "HIGH" else self.med_risk_fill

            ws.cell(row=current_row, column=9, value=str(row.get("anomaly_reasons", "")))
            ws.cell(row=current_row, column=10, value=str(row.get("recommended_action", "")))

            for col_idx in range(1, 11):
                ws.cell(row=current_row, column=col_idx).border = self.cell_border
                ws.cell(row=current_row, column=col_idx).font = self.font_regular
            current_row += 1

        self._auto_adjust_columns(ws)
